"""Structural tests on the committed Phase 1 deliverables."""
import re
import subprocess
import sys
import unittest

import pymupdf
from openpyxl import load_workbook

from phase1_helpers import BUILD, DELIVERABLES, PHASE1, SUPPORTING
import model_spec as m

MIN_PAGES = {
    "P1_01_Project_Proposal.pdf": 3, "P1_02_Business_Requirements.pdf": 3, "P1_03_Business_Questions.pdf": 3,
    "P1_04_Fact_Dimension_Identification.pdf": 5, "P1_05_Grain_Definition.pdf": 3, "P1_06_Star_Schema.pdf": 2,
}


class TestFilesPresent(unittest.TestCase):
    def test_all_deliverables_exist_and_nonempty(self):
        for f in DELIVERABLES + SUPPORTING:
            p = PHASE1 / f
            self.assertTrue(p.is_file(), f)
            self.assertGreater(p.stat().st_size, 1000, f)


class TestPdfs(unittest.TestCase):
    def test_pdfs_open_and_have_text(self):
        for f, n in MIN_PAGES.items():
            p = PHASE1 / f
            self.assertEqual(p.read_bytes()[:5], b"%PDF-", f)
            with pymupdf.open(p) as doc:
                self.assertFalse(doc.is_encrypted, f)
                self.assertFalse(doc.is_repaired, f"{f} has a broken xref table (line-ending corruption?)")
                self.assertGreaterEqual(doc.page_count, n, f)
                for i, page in enumerate(doc):
                    self.assertGreater(len(page.get_text().strip()), 50, f"{f} page {i + 1} is empty")

    def test_documents_carry_project_title(self):
        for f in MIN_PAGES:
            with pymupdf.open(PHASE1 / f) as doc:
                self.assertIn(m.PROJECT, doc[0].get_text(), f)

    def test_schema_page_is_large_format(self):
        with pymupdf.open(PHASE1 / "P1_06_Star_Schema.pdf") as doc:
            self.assertGreaterEqual(doc[0].rect.width, 1600)  # A2 landscape for legibility

    def test_png_diagrams(self):
        pix = pymupdf.Pixmap(str(PHASE1 / "P1_06_Star_Schema.png"))
        self.assertGreaterEqual(pix.width, 3000)
        self.assertGreater(pix.width, pix.height)
        pix = pymupdf.Pixmap(str(PHASE1 / "P1_06_Conceptual_Model.png"))
        self.assertGreaterEqual(pix.width, 2000)


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(PHASE1 / "P1_07_Data_Dictionary.xlsx")

    def test_sheets(self):
        for s in ["Overview", "Tables", "Columns", "Relationships", "Bus_Matrix", "Hierarchies", "Measures",
                  "Lookup_Values", "Business_Questions", "Transformation_Rules", "Source_Profile"]:
            self.assertIn(s, self.wb.sheetnames)

    def test_headers_and_rows(self):
        expect = {"Tables": len(m.TABLES), "Columns": sum(len(t["columns"]) for t in m.TABLES),
                  "Relationships": len(m.relationships()), "Measures": len(m.MEASURES),
                  "Business_Questions": len(m.QUESTIONS)}
        for sheet, n in expect.items():
            ws = self.wb[sheet]
            header = [c.value for c in ws[1]]
            self.assertTrue(all(header), f"{sheet} has a blank header")
            self.assertEqual(ws.max_row - 1, n, sheet)
            self.assertEqual(ws.freeze_panes, "A2", sheet)

    def test_columns_sheet_matches_spec(self):
        ws = self.wb["Columns"]
        hdr = [c.value for c in ws[1]]
        got = {(r[hdr.index("Table")], r[hdr.index("Column")]): (r[hdr.index("Data type (PostgreSQL)")], r[hdr.index("Key role")])
               for r in ws.iter_rows(min_row=2, values_only=True)}
        want = {(t["name"], c[0]): (c[1], c[2]) for t in m.TABLES for c in t["columns"]}
        self.assertEqual(got, want)


class TestDbml(unittest.TestCase):
    def test_refs_point_to_declared_columns(self):
        text = (PHASE1 / "P1_06_Star_Schema.dbml").read_text(encoding="utf-8")
        tables = {}
        current = None
        for line in text.splitlines():
            mt = re.match(r"Table (\w+) ", line)
            if mt:
                current = tables.setdefault(mt.group(1), set())
            elif current is not None and re.match(r"  \w+ \S+ \[", line):
                current.add(line.split()[0])
            elif line.startswith("}"):
                current = None
        self.assertEqual(set(tables), {t["name"] for t in m.TABLES})
        refs = re.findall(r"^Ref: (\w+)\.(\w+) > (\w+)\.(\w+)$", text, re.M)
        self.assertEqual(len(refs), len(m.relationships()))
        for ft, fc, dt, dc in refs:
            self.assertIn(fc, tables[ft])
            self.assertIn(dc, tables[dt])


class TestValidationSuite(unittest.TestCase):
    def test_validate_phase1_passes(self):
        r = subprocess.run([sys.executable, str(BUILD / "validate_phase1.py")], cwd=BUILD, capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        mt = re.search(r"(\d+) checks, (\d+) failed", r.stdout)
        self.assertIsNotNone(mt, r.stdout)
        self.assertEqual(int(mt.group(2)), 0)
        self.assertGreaterEqual(int(mt.group(1)), 269)


if __name__ == "__main__":
    unittest.main()
