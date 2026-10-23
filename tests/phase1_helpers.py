"""Shared helpers for the Phase 1 tests (no raw Olist data required)."""
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PHASE1 = ROOT / "01_Project_Documentation" / "phase1"
BUILD = PHASE1 / "build"
if str(BUILD) not in sys.path:
    sys.path.insert(0, str(BUILD))

DELIVERABLES = [
    "P1_01_Project_Proposal.pdf",
    "P1_02_Business_Requirements.pdf",
    "P1_03_Business_Questions.pdf",
    "P1_04_Fact_Dimension_Identification.pdf",
    "P1_05_Grain_Definition.pdf",
    "P1_06_Star_Schema.pdf",
    "P1_07_Data_Dictionary.xlsx",
]
SUPPORTING = ["P1_06_Star_Schema.png", "P1_06_Conceptual_Model.png", "P1_06_Star_Schema.dbml"]
BUILD_SCRIPTS = ["build_data_dictionary.py", "build_schema_diagram.py", "build_documents.py"]

_FOOTER = re.compile(r"^(Olist Commerce360 \| DL3004 Phase 1 \| .*|Page \d+)$")


def run_script(script, env=None):
    return subprocess.run([sys.executable, str(BUILD / script)], cwd=BUILD, env=env,
                          capture_output=True, text=True)


def pdf_pages_text(path):
    """Text of each page with whitespace collapsed and running footers removed."""
    import pymupdf
    pages = []
    with pymupdf.open(path) as doc:
        for page in doc:
            lines = [ln.strip() for ln in page.get_text().splitlines()]
            lines = [ln for ln in lines if ln and not _FOOTER.match(ln)]
            pages.append(re.sub(r"\s+", " ", " ".join(lines)).strip())
    return pages


def xlsx_values(path):
    """{sheet: [[cell values]]} for a workbook."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}
