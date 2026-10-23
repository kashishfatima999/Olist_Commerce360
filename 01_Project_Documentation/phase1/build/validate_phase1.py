"""Validate the Phase 1 deliverables against model_spec.py and against each other.

Writes 01_Project_Documentation/phase1/build/validation_report.md and exits non-zero on failure.
"""
import re
import sys
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

import model_spec as m

D = Path(__file__).resolve().parents[1]
results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))


def pdf_text(path, pages=None):
    doc = pymupdf.open(path)
    return "\n".join(doc[i].get_text() for i in (pages or range(doc.page_count)))


DELIVERABLES = {
    "P1_01_Project_Proposal.pdf": "Project Proposal PDF",
    "P1_02_Business_Requirements.pdf": "Business Requirements Document PDF",
    "P1_03_Business_Questions.pdf": "10 Business Questions PDF",
    "P1_04_Fact_Dimension_Identification.pdf": "Fact & Dimension Identification PDF",
    "P1_05_Grain_Definition.pdf": "Grain Definition Document PDF",
    "P1_06_Star_Schema.pdf": "Star/Snowflake Schema Diagram (PDF)",
    "P1_06_Star_Schema.png": "Star/Snowflake Schema Diagram (PNG)",
    "P1_07_Data_Dictionary.xlsx": "Data Dictionary Excel",
}

# 1. existence and opening -------------------------------------------------------------
for f, label in DELIVERABLES.items():
    p = D / f
    if not p.exists():
        check(f"{label}: file exists", False, f)
        continue
    try:
        if f.endswith(".pdf"):
            doc = pymupdf.open(p)
            n = doc.page_count
            # layout: every text span inside the page, no replacement glyphs
            overflow, bad = 0, 0
            for pg in doc:
                r = pg.rect
                for b in pg.get_text("dict")["blocks"]:
                    for ln in b.get("lines", []):
                        for sp in ln["spans"]:
                            x0, y0, x1, y1 = sp["bbox"]
                            if x0 < r.x0 - 1 or x1 > r.x1 + 1 or y0 < r.y0 - 1 or y1 > r.y1 + 1:
                                overflow += 1
                            if "�" in sp["text"]:
                                bad += 1
            check(f"{label}: opens", n > 0, f"{n} pages")
            check(f"{label}: no text outside page / no missing glyphs", overflow == 0 and bad == 0,
                  f"overflow spans={overflow}, missing glyphs={bad}")
        elif f.endswith(".png"):
            pix = pymupdf.Pixmap(str(p))
            check(f"{label}: opens", pix.width > 2000, f"{pix.width}x{pix.height}px")
        else:
            wb = load_workbook(p)
            check(f"{label}: opens", True, ", ".join(wb.sheetnames))
    except Exception as e:  # pragma: no cover
        check(f"{label}: opens", False, repr(e))

# 2. model integrity --------------------------------------------------------------------
names = {t["name"] for t in m.TABLES}
for t in m.TABLES:
    pks = [c for c in t["columns"] if c[2] == "PK"]
    check(f"{t['name']}: has primary key", pks, ", ".join(c[0] for c in pks))
    cols = [c[0] for c in t["columns"]]
    check(f"{t['name']}: no duplicate column names", len(cols) == len(set(cols)))
    if t["kind"] == "dimension":
        check(f"{t['name']}: has natural key", any(c[2] == "NK" for c in t["columns"]))
        check(f"{t['name']}: used by at least one fact", any(t["name"] in ds for ds in m.bus_matrix().values()))
    else:
        check(f"{t['name']}: has measures with additivity", all(c[8] for c in t["columns"] if c[2] == "M")
              and any(c[2] == "M" for c in t["columns"]))
        check(f"{t['name']}: grain statement present", t["grain"].startswith("One row per"))
    for c in t["columns"]:
        if c[2] == "FK":
            tt, cc = c[9].split(".")
            check(f"{t['name']}.{c[0]}: FK target exists", tt in names and cc in [x[0] for x in m.TABLE[tt]["columns"]], c[9])
            check(f"{t['name']}.{c[0]}: FK not nullable (unknown member)", not c[3])
            tgt = m.col(tt, cc)
            check(f"{t['name']}.{c[0]}: FK type matches PK type", tgt[1] == c[1], f"{c[1]} vs {tgt[1]}")

# 3. business questions map to valid measures / dimensions / facts -----------------------
measure_names = {x[0] for x in m.MEASURES}
check("At least 10 business questions", len(m.QUESTIONS) >= 10, str(len(m.QUESTIONS)))
req_ids = {r[0] for r in m.REQUIREMENTS}
for q in m.QUESTIONS:
    for f in q[5]:
        check(f"{q[0]}: fact '{f}' exists", f in names and m.TABLE[f]["kind"] == "fact")
    fact_cols = {c[0] for f in q[5] for c in m.TABLE[f]["columns"]}
    for meas in q[3]:
        ok = meas in measure_names or meas in fact_cols
        if meas in measure_names:
            mf = next(x[2] for x in m.MEASURES if x[0] == meas)
            ok = ok and (mf in q[5] or meas in ("Item sales revenue", "Order count"))
        check(f"{q[0]}: measure '{meas}' defined and available in its fact(s)", ok)
    for dtxt in q[4]:
        dn = dtxt.split(" ")[0]
        check(f"{q[0]}: dimension '{dn}' exists and joins to a listed fact", dn in names and
              any(dn in m.bus_matrix()[f] for f in q[5]))
    check(f"{q[0]}: requirement {q[8]} exists", q[8] in req_ids)
for r in m.REQUIREMENTS:
    if r[0] != "BR-12":
        check(f"{r[0]}: covered by a business question", any(q[8] == r[0] for q in m.QUESTIONS))
for x in m.MEASURES:
    check(f"KPI '{x[0]}': fact table exists", x[2] in names)

# 4. data dictionary matches spec -----------------------------------------------------------
wb = load_workbook(D / "P1_07_Data_Dictionary.xlsx")
need = ["Overview", "Tables", "Columns", "Relationships", "Bus_Matrix", "Hierarchies", "Measures", "Lookup_Values",
        "Business_Questions", "Transformation_Rules", "Source_Profile"]
check("Data dictionary: required sheets present", all(s in wb.sheetnames for s in need), ", ".join(wb.sheetnames))
ws = wb["Columns"]
hdr = [c.value for c in ws[1]]
rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
xl_cols = {(r["Table"], r["Column"]): r for r in rows}
spec_cols = {(t["name"], c[0]): c for t in m.TABLES for c in t["columns"]}
check("Data dictionary: same columns as model", set(xl_cols) == set(spec_cols),
      f"{len(xl_cols)} in xlsx, {len(spec_cols)} in model")
check("Data dictionary: every column has type, key role, description, source, rule",
      all(r["Data type (PostgreSQL)"] and r["Key role"] and r["Description"] and r["Source"] and r["Transformation rule"]
          for r in rows))
check("Data dictionary: types/roles identical to model",
      all(xl_cols[k]["Data type (PostgreSQL)"] == v[1] and xl_cols[k]["Key role"] == v[2] for k, v in spec_cols.items()
          if k in xl_cols))
rel_rows = list(wb["Relationships"].iter_rows(min_row=2, values_only=True))
check("Data dictionary: relationships match model", len(rel_rows) == len(m.relationships()),
      f"{len(rel_rows)} relationships")

# 5. diagram matches dictionary -------------------------------------------------------------
diag = pdf_text(D / "P1_06_Star_Schema.pdf", [0])
missing = [f"{t}.{c}" for (t, c) in spec_cols if c not in diag]
missing_tables = [t["name"] for t in m.TABLES if f"{m.SCHEMA}.{t['name']}" not in diag]
check("Schema diagram: every table shown", not missing_tables, ", ".join(missing_tables))
check("Schema diagram: every column of the dictionary shown", not missing, ", ".join(missing[:10]))
check("Schema diagram: every grain statement shown",
      all(t["grain"].split(".")[0][:40] in diag.replace("\n", " ") for t in m.TABLES))
dbml = (D / "P1_06_Star_Schema.dbml").read_text(encoding="utf-8")
check("DBML: all tables and refs", all(f"Table {t['name']} " in dbml for t in m.TABLES)
      and dbml.count("\nRef: ") == len(m.relationships()))

# 6. documents reference the same names --------------------------------------------------------
docs = {f: pdf_text(D / f).replace("\n", " ") for f in DELIVERABLES if f.endswith(".pdf") and f.startswith("P1_0") and "06" not in f}
for f in ("P1_04_Fact_Dimension_Identification.pdf", "P1_05_Grain_Definition.pdf"):
    absent = [t["name"] for t in m.TABLES if t["name"] not in docs[f]]
    check(f"{f}: mentions every table", not absent, ", ".join(absent))
g = re.sub(r"\s+", " ", docs["P1_05_Grain_Definition.pdf"])
check("P1_05: every fact grain statement verbatim", all(re.sub(r"\s+", " ", t["grain"]) in g for t in m.FACTS))
q = docs["P1_03_Business_Questions.pdf"]
check("P1_03: all question IDs present", all(x[0] in q for x in m.QUESTIONS))
p4 = docs["P1_04_Fact_Dimension_Identification.pdf"]
for term in ["Conceptual model", "Logical model", "star schema", "Hierarchies", "Surrogate key", "Natural key",
             "Accumulating snapshot", "Transaction", "Additive", "Non-additive", "bus matrix"]:
    check(f"P1_04: covers '{term}'", term.lower() in p4.lower())
# forbidden claims: phase 1 must not claim implementation or analysis results
alltext = " ".join(docs.values()).lower()
for bad in ["has been implemented", "was loaded", "dashboard shows", "we found that revenue", "increased by"]:
    check(f"No unsupported claim '{bad}'", bad not in alltext)
for bad in ["profit margin is", "inventory level"]:
    check(f"No unsupported metric '{bad}'", bad not in alltext)

# report ---------------------------------------------------------------------------------------
fails = [r for r in results if not r[1]]
lines = ["# Phase 1 Validation Report", "",
         f"Generated by `build/validate_phase1.py`. Checks: {len(results)}, passed: {len(results) - len(fails)}, "
         f"failed: {len(fails)}.", "", "| Check | Result | Detail |", "|---|---|---|"]
lines += [f"| {n} | {'PASS' if ok else 'FAIL'} | {d} |" for n, ok, d in results]
(Path(__file__).parent / "validation_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"{len(results)} checks, {len(fails)} failed")
for n, ok, d in fails:
    print("FAIL:", n, "|", d)
sys.exit(1 if fails else 0)
