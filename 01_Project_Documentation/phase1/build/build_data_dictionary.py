"""Build P1_07_Data_Dictionary.xlsx from model_spec.py."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import model_spec as m

OUT = Path(__file__).resolve().parents[1] / "P1_07_Data_Dictionary.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FACT_FILL = PatternFill("solid", fgColor="FCE4D6")
DIM_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

SOURCE_PROFILE = [  # file, rows, columns, business key, key unique?, used for
    ("olist_orders_dataset.csv", 99441, 8, "order_id", "Yes", "fact_orders, keys/dates for all facts"),
    ("olist_order_items_dataset.csv", 112650, 7, "order_id + order_item_id", "Yes", "fact_order_items; aggregates in fact_orders"),
    ("olist_order_payments_dataset.csv", 103886, 5, "order_id + payment_sequential", "Yes", "fact_payments, dim_payment_type; aggregates in fact_orders"),
    ("olist_order_reviews_dataset.csv", 99224, 7, "review_id + order_id (review_id alone not unique: 814 dup.)", "Yes (composite)", "fact_reviews; aggregates in fact_orders"),
    ("olist_customers_dataset.csv", 99441, 5, "customer_id (per order); customer_unique_id (person, 96,096 distinct)", "customer_id only", "dim_customer"),
    ("olist_products_dataset.csv", 32951, 9, "product_id", "Yes", "dim_product"),
    ("olist_sellers_dataset.csv", 3095, 4, "seller_id", "Yes", "dim_seller"),
    ("product_category_name_translation.csv", 71, 2, "product_category_name", "Yes", "dim_product (English category)"),
    ("olist_geolocation_dataset.csv", 1000163, 5, "none (19,015 zip prefixes, 52.6 rows each)", "No", "Latitude/longitude in dim_customer, dim_seller"),
]

RULES = [  # id, area, rule, reason / evidence
    ("TR-01", "All IDs", "Trim whitespace; keep 32-char hexadecimal IDs as VARCHAR(32).", "IDs are hash strings."),
    ("TR-02", "Zip prefixes", "Read as text and left-pad to 5 digits.", "Source stores '01037'; reading as integer drops the leading zero (23,995 customer rows)."),
    ("TR-03", "City names", "Trim, remove accents consistently, convert to title case.", "Geolocation has 8,011 city spellings for 27 states."),
    ("TR-04", "Timestamps", "Cast text to TIMESTAMP; date_key = date part as YYYYMMDD; NULL -> -1.", "Approval (160), carrier (1,783) and delivery (2,965) timestamps are NULL for undelivered orders."),
    ("TR-05", "Customer", "Map orders.customer_id -> customers.customer_unique_id -> dim_customer.customer_key.", "customer_id is created per order; customer_unique_id identifies the person (DEC-007)."),
    ("TR-06", "Customer location", "Take zip/city/state from the person's most recent order (SCD Type 1).", "250 persons have more than one zip prefix (39 more than one state)."),
    ("TR-07", "Product category", "NULL category -> 'sem_categoria' / 'uncategorized'; map 2 untranslated names manually.", "610 products lack a category; 'pc_gamer' and 'portateis_cozinha_e_preparadores_de_alimentos' are not in the translation file."),
    ("TR-08", "Product columns", "Rename product_name_lenght / product_description_lenght to *_length.", "Source spelling error."),
    ("TR-09", "Geolocation", "Keep points with latitude -34..6 and longitude -74..-34; take median per zip prefix.", "261,831 fully duplicated rows; coordinates up to lat 45 / lng 121 exist (outside Brazil)."),
    ("TR-10", "Unknown members", "Insert key -1 'Unknown' row in every dimension; unmatched lookups map to -1.", "Guarantees no NULL foreign keys (DEC-014)."),
    ("TR-11", "Payments", "payment_installments = 0 -> 1.", "2 rows have 0 instalments, which is not a valid payment plan."),
    ("TR-12", "Delivery measures", "Compute lags from timestamps in days (2 decimals); delay in whole days from date parts; is_late only when delivered.", "Estimated delivery date has no time component (all 00:00)."),
    ("TR-13", "Timestamp sequence", "Flag (do not delete) orders where carrier date < approval date (1,359) or delivery < carrier (23); exclude negative lags from stage averages.", "Source data quality issue; keeps row counts reconcilable."),
    ("TR-14", "Order aggregates", "item_count, order_item_value, order_freight_value, order_payment_value, review_count, avg_review_score are aggregated per order_id BEFORE joining to orders.", "Prevents fan-out: a naive items x payments join inflates payment totals by BRL 4.30 M."),
    ("TR-15", "Translation file", "Read with UTF-8-SIG encoding.", "File starts with a byte-order mark."),
    ("TR-16", "Reconciliation (Phase 2)", "Row counts and SUM(price), SUM(freight_value), SUM(payment_value) in facts must equal source totals.", "Source totals: price 13,591,643.70; freight 2,251,909.54; payment 16,008,872.12."),
]


def write_sheet(ws, header, rows, widths, fills=None):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, WRAP, BORDER
    for i, r in enumerate(rows):
        ws.append(list(r))
        for c in ws[ws.max_row]:
            c.alignment, c.border = WRAP, BORDER
            if fills and fills[i]:
                c.fill = fills[i]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def fill_for(table):
    return FACT_FILL if m.TABLE[table]["kind"] == "fact" else DIM_FILL


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    info = [
        ("Project", f"{m.PROJECT} - {m.SUBTITLE}"),
        ("Course", m.COURSE),
        ("Deliverable", "Phase 1 - Data Dictionary (P1_07)"),
        ("Target schema", f"PostgreSQL schema '{m.SCHEMA}' (staging: project_staging, reporting: project_reporting)"),
        ("Model", "Star schema: 4 fact tables sharing 6 conformed dimensions (fact constellation)."),
        ("Source", "Kaggle 'Brazilian E-Commerce Public Dataset by Olist' (olistbr/brazilian-ecommerce, version 2), 9 CSV files."),
        ("Status", "Design (Phase 1). Tables are not yet implemented; row counts are expected values from source profiling."),
        ("Version / date", "1.0 - 2026-10-25"),
        ("Prepared by", "[Team member names - Student IDs]"),
        ("", ""),
        ("Key role legend", ""),
        ("PK", "Primary key of the warehouse table"),
        ("NK", "Natural (business) key from the source system; unique in the dimension"),
        ("FK", "Foreign key to a dimension surrogate key"),
        ("DD", "Degenerate dimension (identifier/attribute kept in the fact without its own dimension)"),
        ("M", "Measure (numeric fact); see 'Additivity' column"),
        ("A", "Descriptive attribute"),
        ("", ""),
        ("Sheets", "Tables, Columns, Relationships, Bus_Matrix, Hierarchies, Measures, Lookup_Values, "
                   "Business_Questions, Transformation_Rules, Source_Profile"),
    ]
    for r in info:
        ws.append(r)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = WRAP
    ws["A1"].font = Font(bold=True, size=13)

    # Tables
    rows, fills = [], []
    for t in m.TABLES:
        rows.append([t["name"], t["kind"].title(), t["type"], t.get("process", "-"), t["grain"],
                     t["natural_key"], t["surrogate_key"], t["source"], t["scd"] or "-", t["est_rows"],
                     "; ".join(t.get("hierarchies", [])) or "-", len(t["columns"])])
        fills.append(fill_for(t["name"]))
    write_sheet(wb.create_sheet("Tables"),
                ["Table", "Kind", "Table type", "Business process", "Grain", "Natural / business key",
                 "Surrogate key", "Source", "SCD", "Expected rows", "Hierarchies", "Columns"],
                rows, [20, 11, 26, 22, 55, 26, 24, 45, 22, 16, 32, 9], fills)

    # Columns
    rows, fills = [], []
    for t in m.TABLES:
        for i, c in enumerate(m.TABLE[t["name"]]["columns"], 1):
            name, typ, role, nullable, desc, src, rule, ex, add, ref = c
            rows.append([t["name"], i, name, typ, role, "NULL" if nullable else "NOT NULL", ref or "-",
                         desc, src, rule, ex, add or "-"])
            fills.append(fill_for(t["name"]))
    write_sheet(wb.create_sheet("Columns"),
                ["Table", "#", "Column", "Data type (PostgreSQL)", "Key role", "Nullable", "References",
                 "Description", "Source", "Transformation rule", "Example", "Additivity"],
                rows, [18, 4, 28, 16, 8, 10, 28, 50, 36, 46, 34, 13], fills)

    # Relationships
    rows = [[f, fk, d, pk, "Many-to-one (*:1)", "Mandatory (unknown member -1 when missing)", role]
            for f, fk, d, pk, role in m.relationships()]
    write_sheet(wb.create_sheet("Relationships"),
                ["Fact table", "FK column", "Dimension", "Dimension PK", "Cardinality", "Optionality", "Role / meaning"],
                rows, [18, 28, 18, 18, 18, 34, 60])

    # Bus matrix
    bm = m.bus_matrix()
    dims = [d["name"] for d in m.DIMS]
    rows = []
    for f in m.FACTS:
        rows.append([f["name"], f["process"], f["type"]] +
                    [("X (" + str(len(bm[f["name"]][d])) + " roles)" if len(bm[f["name"]].get(d, [])) > 1 else "X")
                     if d in bm[f["name"]] else "" for d in dims])
    write_sheet(wb.create_sheet("Bus_Matrix"), ["Fact table", "Business process", "Fact type"] + dims,
                rows, [18, 28, 26] + [16] * len(dims))

    # Hierarchies
    rows = [[d["name"], h, d["natural_key"]] for d in m.DIMS for h in d["hierarchies"]]
    write_sheet(wb.create_sheet("Hierarchies"), ["Dimension", "Hierarchy (top > bottom)", "Lowest level natural key"],
                rows, [20, 50, 30])

    # Measures
    write_sheet(wb.create_sheet("Measures"),
                ["Business measure / KPI", "Definition", "Fact table", "Calculation", "Additivity", "Note"],
                m.MEASURES, [28, 45, 18, 60, 18, 50])

    # Lookups
    rows = [["dim_order_status", code, name, grp, "TRUE" if fin else "FALSE"] for code, name, grp, fin in m.ORDER_STATUS]
    rows += [["dim_payment_type", code, name, grp, ""] for code, name, grp in m.PAYMENT_TYPES]
    rows += [["state -> region", uf, nm, rg, ""] for uf, (nm, rg) in m.STATES.items()]
    write_sheet(wb.create_sheet("Lookup_Values"), ["Lookup", "Code", "Name", "Group / region", "Final status"],
                rows, [18, 16, 24, 16, 12])

    # Business questions
    rows = [[q[0], q[1], q[2], "; ".join(q[3]), "; ".join(q[4]), "; ".join(q[5]), q[6], q[7], q[8]] for q in m.QUESTIONS]
    write_sheet(wb.create_sheet("Business_Questions"),
                ["ID", "Question", "Purpose", "Measures", "Dimensions", "Fact table(s)", "Filters / period",
                 "Limitation", "Requirement"],
                rows, [7, 50, 34, 30, 30, 22, 30, 40, 11])

    write_sheet(wb.create_sheet("Transformation_Rules"), ["ID", "Area", "Rule", "Reason / evidence"],
                RULES, [7, 18, 70, 70])

    write_sheet(wb.create_sheet("Source_Profile"),
                ["Source file", "Rows", "Columns", "Business key", "Key unique?", "Used for"],
                SOURCE_PROFILE, [40, 11, 9, 50, 16, 50])
    for r in wb["Source_Profile"].iter_rows(min_row=2, min_col=2, max_col=2):
        r[0].number_format = "#,##0"

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
